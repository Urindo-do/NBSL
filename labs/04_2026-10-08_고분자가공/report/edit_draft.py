"""팀원 docx 초안(t.docx)에 수정 요청·부족분을 반영한다.
python3 edit.py <in.docx> <out.docx> <images_dir> [--blank-cover]
--blank-cover: 저장소용. 표지 조원 이름과 문서 속성 작성자를 비운다.
"""
import copy, sys, re
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn

src, dst, IMG = sys.argv[1:4]
BLANK = '--blank-cover' in sys.argv
d = Document(src)
P = list(d.paragraphs)          # 원래 번호 고정 (삽입 전 기준)
T = d.tables


def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def replace_in(p, old, new):
    t = p.text
    assert t.count(old) == 1, (old, t[:60])
    set_text(p, t.replace(old, new))


def add_after(anchor_el, text, tmpl=None):
    """anchor_el 뒤에 본문 문단(tmpl 서식 복제)을 넣고 그 요소를 돌려준다."""
    tmpl = tmpl if tmpl is not None else P[92]
    new = copy.deepcopy(tmpl._p)
    anchor_el.addnext(new)
    from docx.text.paragraph import Paragraph
    para = Paragraph(new, tmpl._parent)
    set_text(para, text)
    return new


_docpr = [9000]


def img_para_after(anchor_el, files, h_cm):
    """P[93](사진 2장 문단)을 복제해 사진을 바꿔 넣는다."""
    from PIL import Image
    new = copy.deepcopy(P[93]._p)
    drawings = new.findall('.//' + qn('w:drawing'))
    assert len(drawings) == len(files) == 2
    for dr, f in zip(drawings, files):
        path = f'{IMG}/{f}'
        rid, _ = d.part.get_or_add_image(path)
        w, h = Image.open(path).size
        cy = int(Cm(h_cm)); cx = int(cy * w / h)
        for ext in dr.iter(qn('wp:extent')):
            ext.set('cx', str(cx)); ext.set('cy', str(cy))
        for ext in dr.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ext'):
            ext.set('cx', str(cx)); ext.set('cy', str(cy))
        for blip in dr.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            blip.set(qn('r:embed'), rid)
        for dp in dr.iter(qn('wp:docPr')):
            _docpr[0] += 1; dp.set('id', str(_docpr[0])); dp.set('name', f'Picture {_docpr[0]}'); dp.set('descr', f)
        for cp in dr.iter('{http://schemas.openxmlformats.org/drawingml/2006/picture}cNvPr'):
            cp.set('id', str(_docpr[0])); cp.set('name', f)
    anchor_el.addnext(new)
    return new


def cell_set(tbl, r, c, text):
    p = tbl.rows[r].cells[c].paragraphs[0]
    set_text(p, text)

# ---------------- 표지 ----------------
if BLANK:
    for i in range(9, 13):
        t = P[i].text
        set_text(P[i], t.split(':')[0] + ': ')
    cp = d.core_properties
    cp.author = ''; cp.last_modified_by = ''

# ---------------- 1. 서론 ----------------
replace_in(P[48], '원판의 표면 패턴을 복제하고,', '원판의 표면 형상을 복제하는 공정을 익히고')
set_text(P[50],
    '실리콘 웨이퍼 원판에 액상 수지를 도포하고 PET 필름을 밀착시킨 상태에서 UV로 경화하였다. '
    '패턴이 새겨진 웨이퍼는 한 장에 200만~300만 원 정도로 고가여서 이번에는 패턴이 없는 플랫(flat) 웨이퍼를 원판으로 사용하였다. '
    '따라서 제작한 시편은 무늬가 없는 평판 PUA 필름이다. 패턴 복제의 원리는 함께 소개된 패턴 필름의 관찰과 문헌으로 정리하였다. '
    'PET는 얇은 수지층을 지지하고 탈형 후 시편을 취급하는 데 사용하였다. '
    '제작 과정과 시편의 외관을 중심으로 관찰하였으며, 나노구조의 치수와 복제 정확도는 별도의 표면 측정이 필요한 평가 항목이다.')

# ---------------- 2. 재료와 원리 ----------------
cell_set(T[0], 3, 0, '원판(실리콘 웨이퍼)')
cell_set(T[0], 3, 2, '패턴 없는 반원형 플랫 원판 사용·가장자리 기존 파손')
set_text(P[77], P[77].text + ' 수업에서는 아세톤과 에탄올로 세척한 웨이퍼를 오븐에서 5분 정도 말린다고 안내하였다. 촬영한 건조기가 이 단계에 쓰인 장비로 보인다.')

set_text(P[84],
    'PDMS(폴리디메틸실록산)는 실리콘 계열 고분자이다. 대표 상용 제품인 SYLGARD 184는 베이스와 경화제를 무게비 10 : 1로 섞어 실온 또는 가열로 경화하며 '
    '경화 온도가 높을수록 인장강도와 경도, 영률이 커진다 [5]. 같은 자료는 이 제품이 소수성 PDMS 막을 만들고 미세유체 칩 같은 소프트 리소그래피에 쓰인다고 설명한다 [5]. '
    '수업에서도 베이스와 경화제를 10 : 1로 섞어 틀에 붓고 열로 굳혀 PDMS 플레이트를 만든다고 소개하였다. '
    '그림 2-3(a)의 용기 라벨에는 Silicone Elastomer Base와 경화제를 무게비로 섞으라는 안내가 있으나 소분 판매 제품이라 원제품명은 확인하지 못하였다. '
    '경화된 PDMS는 패턴 복제용 몰드로도 쓰이며 PUA를 PDMS 몰드에 떨어뜨려 형상을 복제한 연구가 있다 [1]. '
    '이번 실습에서 PDMS는 재료와 원리만 다루었고 직접 제작하지는 않았다. PUA의 UV 경화와 PDMS의 열경화는 서로 다른 제작 과정으로 구분하였다.')
add_after(P[85]._p,
    'PCL(폴리카프로락톤)과 PLGA(폴리락트산-글리콜산 공중합체)는 조직공학 지지체에 쓰이는 생분해성 폴리에스터이다 [6]. '
    '두 재료는 디클로로메탄(DCM) 같은 용매에 녹인 뒤 전기방사하거나 용액을 몰드에 붓고 용매를 증발시키는 용매 캐스팅으로 지지체를 만든다 [6]. '
    '수업에서도 PCL과 PLGA를 DCM에 녹여 교반한 뒤 압력이나 열을 가해 넓은 필름으로 만든다고 설명하였다. '
    'PCL은 유리전이온도가 약 −60 °C, 녹는점이 59–64 °C로 낮아 상온에서 유연하지만 소수성이어서 표면에 세포가 덜 붙는다 [7]. '
    '수업에서는 PCL이 부드러워 패치로 다루기 편하고 PLGA는 더 단단하지만 깨지기 쉽다고 소개하였다. 그림 2-3(b)의 시편은 이번에 제작하지 않고 실물로만 살펴보았다.')

# ---------------- 3. 실험 방법 ----------------
set_text(P[92],
    '먼저 PET 필름을 원판보다 2~3배 넓게 가위로 절단하였다. PET가 작으면 수지가 필름 밖으로 흘러나와 원판과 작업대에 붙기 때문이다. '
    '자르기 전에는 PET를 빛에 비춰 흠집과 지문이 없는 깨끗한 부분을 골랐다. '
    'PET가 접히거나 들뜨지 않도록 준비하는 것은 원판과의 균일한 밀착에 중요하다.')
# 그림 3-2에 (c)(d) 추가
e = img_para_after(P[94]._p, ['04_prep_wash-bottles.jpg', '04_prep_air-blower.jpg'], 4.6)
add_after(e, '(c) 세척에 쓴 에탄올 세척병     (d) 에어 블로어로 PET 먼지 제거', P[94])
set_text(P[95], '그림 3-2 PET 절단, 복제에 사용한 원판과 세척 준비')
set_text(P[96],
    '수업에서는 원판을 아세톤과 에탄올로 세척하도록 안내하였다. 작업대에서는 에탄올 세척병을 확인하였다(그림 3-2(c)). '
    '세척한 원판은 건조하였고 PET 표면의 먼지는 고무 에어 블로어로 불어 날린 뒤 깨끗한 면에 먼지가 붙지 않게 내려놓았다(그림 3-2(d)). '
    '패턴 원판은 지문과 오염, 손상을 막기 위해 장갑을 끼고 가장자리를 잡아 다룬다. '
    '세척과 건조는 수지가 접촉할 면의 오염과 잔류액을 줄여 균일한 밀착을 확보하기 위한 과정이다. '
    '세척제는 원판의 재질과 표면 처리에 적합한 것을 사용해야 하며, 건조 후에는 남은 액체와 오염을 확인하는 것이 중요하다.')

set_text(P[99],
    '준비한 원판에 MINS-311RM 수지를 도포한 뒤 PET 필름을 접촉시켰다. 수지는 100–1000 µL 피펫으로 병에서 바로 빨아들여 원판 위 여러 지점에 나누어 떨어뜨렸다(그림 3-3). '
    '수지가 원판과 PET 사이에서 펼쳐지도록 하여 연속적인 층을 형성하고자 하였다. '
    '도포량이 너무 적으면 일부 영역에 수지가 부족할 수 있고, 너무 많으면 가장자리로 흘러나오거나 두께가 불균일해질 수 있다.')
set_text(P[100],
    '도포량, 분주 도구와 도포 횟수는 수지층의 두께와 균일성에 영향을 줄 수 있는 기록 항목이다. '
    '이번 실습에서 분주 도구는 피펫으로 확인하였으나 도포량과 방울 수는 정량으로 기록하지 않았다. '
    '반복 제작 시에는 사용량과 도포 위치를 기록하여 시편 간 차이를 비교할 필요가 있다.')
set_text(P[103],
    'PET 필름의 깨끗한 면을 수지 쪽으로 향하게 덮고 수지가 스스로 퍼지도록 잠시 기다린 뒤 롤러로 밀어 수지를 펼치면서 기포를 제거하였다. '
    '수업에서는 롤러를 끊지 말고 이어서 밀라고 안내하였다. 민 뒤에는 빛에 비춰 기포나 하얗게 덜 퍼진 부분이 남았는지 확인한다. '
    '롤러는 필름과 원판의 접촉을 확보하고 두 표면 사이에 갇힌 공기를 줄이는 역할을 한다. '
    '국부적인 기포나 이물질은 균일한 밀착을 방해할 수 있으므로 밀착 상태를 확인하는 것이 중요하다.')
# 그림 3-3(새): 3.3절 피펫 도포 사진. 기존 3-3 → 3-4, 3-4 → 3-5
e = img_para_after(P[99]._p, ['05_pua_pipette-draw.jpg', '05_pua_drop-on-wafer.jpg'], 4.0)
e = add_after(e, '(a) 피펫으로 311RM 취하기     (b) 원판 위에 수지 떨어뜨리기', P[94])
add_after(e, '그림 3-3 피펫을 이용한 311RM 도포', P[95])
set_text(P[107], P[107].text.replace('그림 3-3', '그림 3-4'))
set_text(P[117], P[117].text.replace('그림 3-4', '그림 3-5'))

set_text(P[111],
    '수지와 PET 필름이 원판에 밀착된 상태를 유지하여 UV 경화기에 넣고 빛을 조사하였다. '
    '원판 한쪽이 들리면 PET와 수지 사이로 공기가 들어갈 수 있어 수평을 유지한 채 서랍에 넣었다. '
    '시편은 경화기에 약 5~10분 넣어 두었다가 바로 꺼냈다. 서랍 안에는 여러 조의 시편이 받침 종이째 놓여 있었다. '
    '이 과정에서 원판과 접촉한 수지의 형상을 고정한 뒤 탈형을 진행하였다. 조사 조건을 재현하려면 사용한 광원, 시편 위치, 조사 세기와 시간을 함께 기록해야 한다.')
set_text(P[112],
    '실습 당시 UV 경화기의 타이머에서 40.0을 확인하였다. 표시부에는 단위 표기가 없고 시작·모드 버튼만 있어 이 값이 초인지 분인지, '
    '실제로 넣어 둔 5~10분과 어떤 관계인지는 확인하지 못하였다. '
    '이후에는 장비의 설정과 조사 시작·종료 시점을 함께 기록하여 시간 조건을 명확히 관리할 필요가 있다.')
add_after(P[112]._p,
    '참고로 같은 MINS-311RM을 실리콘 원판에 떨어뜨리고 PET 필름 너머로 365 nm UV를 30초 조사해 경화한 뒤 PET를 떼어 낸 보고가 있다 [8]. '
    '필요한 시간은 광원의 세기와 시편 위치에 따라 달라지므로 이 값은 이번 실습 조건이 아니라 비교를 위한 참고값이다. '
    '2.3절의 관계에 따르면 같은 조사량을 얻는 데 조사 세기가 약한 장비일수록 더 긴 시간이 필요하다.')

set_text(P[114],
    '경화기에서 꺼낸 뒤 수지가 굳었는지 먼저 확인하고 PET 필름을 원판에서 들어 올려 분리하였다(그림 3-5). '
    '경화가 덜 된 상태로 떼면 수지가 액체 그대로 묻어 나와 시편이 만들어지지 않는다. '
    '탈형할 때에는 가장자리에서부터 서서히 분리하여 수지층과 표면 구조에 가해지는 급격한 힘을 줄이는 것이 중요하다. '
    '또한 원판에 남은 수지와 시편의 찢어짐 또는 균열을 확인하여 탈형 상태를 평가할 필요가 있다.')
add_after(P[118]._p,
    '수업에서는 떼어 낸 필름을 하루 정도 더 두어 완전히 경화해야 변형이 생기지 않는다고 안내하였다. '
    '이번 기록에서는 이 후경화를 어떤 방식으로 했는지 확인하지 못하였다. '
    '문헌에서는 경화가 끝난 PUA에 남은 라디칼을 없애려고 UV를 3시간 더 조사한 예가 있다 [1].')

# ---------------- 4. 결과 ----------------
cell_set(T[1], 3, 1, '반원형 플랫 원판(패턴 없음), 가장자리 일부 파손')
# UV 경화 행 추가 (밀착 행 뒤)
row = copy.deepcopy(T[1].rows[4]._tr)
T[1].rows[4]._tr.addnext(row)
cell_set(T[1], 5, 0, 'UV 경화')
cell_set(T[1], 5, 1, '경화기에 약 5~10분 넣은 뒤 꺼냄')
cell_set(T[1], 5, 2, '타이머 40.0의 단위 미확인')
replace_in(P[129], '원판의 패턴 치수에 대한', '원판 상태에 대한')
replace_in(P[137], '실제 조사 시간도 충분히 기록하지 않아,', '조사 시간도 약 5~10분으로 대략적으로만 남아')

set_text(P[139],
    'UV 경화 후 투명 필름을 원판에서 들어 올려 분리하였다. 제작한 필름은 투명하고 무늬가 없는 평판이었으며 '
    '원판이 플랫이었으므로 보는 각도에 따라 색이 달라지는 현상은 나타나지 않았다(그림 4-3(a)). '
    '떼어 낸 뒤 받침 종이에는 원판 둘레로 번진 PUA 자국이 남아 있었다. 종이를 깔지 않았다면 이 수지가 작업대에 굳어 붙었을 것이다. '
    '롤러 밀착 과정에서는 실습 전부터 파손되어 있던 원판 가장자리 부위의 기포를 완전히 제거하기 어려웠다. '
    '이 경험으로 원판의 표면 상태와 가장자리의 지지 상태도 밀착 과정에서 고려해야 함을 알 수 있었다.')
set_text(P[140],
    '실습 중 비교용으로 살펴본 패턴 필름은 사각 영역이 보는 각도에 따라 여러 색으로 반사되었다(그림 4-3(b)). '
    '이 필름은 패턴 웨이퍼로 만든 것으로 우리 조가 제작한 시편이 아니다. 촬영한 사진에서는 필름 위 손글씨가 N 600nm로 읽힌다. '
    '현장에서 800 nm로 보이는 표기도 보았으나 사진으로 남은 것은 600 nm 필름이다. '
    '표기가 주기·폭·높이 중 어느 치수를 뜻하는지는 확인하지 못하였다. 따라서 색 반사는 패턴 필름의 광학적 특징을 이해하는 참고 관찰로 정리하였다. '
    '패턴 시편의 나노구조 규격과 복제 품질은 원판 사양 및 SEM·AFM 측정으로 평가할 필요가 있다.')
e = img_para_after(P[140]._p, ['05_pua_flat-film.jpg', '06_pua_pattern-600nm.jpg'], 4.3)
e = add_after(e, '(a) 제작한 플랫 PUA 필름     (b) 비교용 패턴 필름(N 600nm 표기)', P[135])
add_after(e, '그림 4-3 제작한 시편과 비교용 패턴 필름', P[136])

# ---------------- 5. 고찰 ----------------
set_text(P[152],
    '이번 실습에서는 시편을 UV 경화기에 약 5~10분 넣어 두었다가 바로 꺼냈다. 다만 정확한 시간과 타이머 표시 40.0의 단위를 기록하지 않았고 조사 세기도 측정하지 않아 '
    '이 시간을 최적 경화 시간으로 제시하기는 어렵다. 문헌에서는 같은 311RM을 365 nm에서 30초 조사해 경화하였으므로 [8] 장비의 세기에 따라 필요한 시간이 크게 다를 수 있다. '
    '최적 조건을 정하려면 나머지 조건을 일정하게 유지하고 조사 시간을 달리하여 경화 상태, 탈형 상태와 패턴 형상을 비교해야 한다.')
set_text(P[155],
    '301RM과 311RM은 각각 soft PUA와 hard PUA로 구분된다 [1, 9]. 문헌에 보고된 경화 후 영률은 311RM이 약 320 MPa, 301RM이 약 20 MPa이고 '
    '끊어질 때까지 늘어나는 비율은 각각 9 %, 45 %이다 [1]. 311RM은 형상을 잘 유지하지만 덜 늘어나고 301RM은 유연한 대신 형상이 쉽게 변할 수 있다. '
    '수지 등급이 탈형성에 미치는 영향을 평가하려면 같은 원판과 제작 조건에서 두 수지를 비교하고, 복제 구조와 기계적 응답을 측정해야 한다. '
    '이번 실습에서는 311RM을 사용하였으므로 등급 간 우열을 판단하기 어렵다. 비교 실험에서는 원판의 표면 상태와 분리 방법도 일정하게 유지해야 한다.')
set_text(P[164],
    '이번 실습에서는 제작 조건을 충분히 수치로 기록하지 않았고 표면 형상도 측정하지 않아 공정 재현성과 복제 정확도를 정량적으로 평가하기 어려웠다. '
    '플랫 원판으로 제작하여 패턴 복제 품질을 직접 확인할 시편이 없었다. UV 경화 시간은 약 5~10분으로만 남았고 후경화 방법도 확인하지 못하였다. '
    '또한 원판 가장자리의 기존 파손으로 기포 제거에 어려움을 겪었다. 이후에는 원판의 상태와 작업 범위를 먼저 확인하고 PET 규격, 도포량, 건조 조건, '
    'UV 파장·조사 세기·실제 시간을 함께 기록할 필요가 있다. 기포 제거 전후와 탈형 전후의 동일한 영역을 촬영하면 공정 중 발생한 변화를 비교할 수 있을 것이다.')
add_after(P[168]._p,
    '생체재료로 쓸 고분자를 고를 때에는 기계적 성질과 함께 분해성과 표면 성질을 고려해야 한다. '
    'PCL과 PLGA는 생분해성 폴리에스터여서 조직공학 지지체로 쓰이며 [6], PCL은 소수성이어서 표면에 세포가 덜 붙는다 [7]. '
    '같은 생분해성 재료라도 세포가 붙어 자라야 하는 용도라면 표면 성질을 함께 조절할 필요가 있다. '
    '이번 실습에서는 세포 실험을 하지 않아 제작한 PUA 필름과 소개된 재료의 생체적합성은 평가하지 않았다.')

# ---------------- 6. 결론 ----------------
replace_in(P[170], '원판의 표면 형상을 복제하는', '패턴 없는 플랫 원판의 표면 형상을 복제하는')

# ---------------- 7. 참고문헌 ----------------
set_text(P[176], '[4] Physics LibreTexts (UC Davis Physics 9B). 3.3: Diffraction Gratings. '
         'https://phys.libretexts.org/Courses/University_of_California_Davis/UCD:_Physics_9B__Waves_Sound_Optics_Thermodynamics_and_Fluids/03:_Physical_Optics/3.03:_Diffraction_Gratings (열람 2026.10.13)')
REFS = [
    '[5] MilliporeSigma. SYLGARD® 184 Silicone Elastomer Kit (Product No. 761036). https://www.sigmaaldrich.com/product/aldrich/761036 (열람 2026.10.13)',
    '[6] Rosalia, M. et al. Investigation on Electrospun and Solvent-Casted PCL-PLGA Blends Scaffolds Embedded with Induced Pluripotent Stem Cells for Tissue Engineering. '
    'Pharmaceutics, 2023, 15(12), 2736. DOI: 10.3390/pharmaceutics15122736.',
    '[7] Dolatkhahan, M.; Peiravi, M.; Sherafat, Z. Morphology and Mechanical Properties of 3D Porous PLA/PCL Scaffolds. Iranian Journal of Materials Forming, 2024, 11(3), 21–27. '
    'https://ijmf.shirazu.ac.ir/article_7767_7a5e5afabc0729125185e581023f9bf5.pdf',
    '[8] Jang, S.; Jeon, H.-J.; An, C. J.; Jin, M. L.; Jung, H.-T. 10 nm scale nanopatterning on flexible substrates by secondary sputtering phenomenon and their applications '
    'in high performance, flexible and transparent conducting films. J. Mater. Chem. C, 2014, Supporting Information. https://www.rsc.org/suppdata/tc/c4/c4tc00012a/c4tc00012a1.pdf',
    '[9] Kim, S. M. et al. Repeated Shape Recovery of Clustered Nanopillars by Mechanical Pulling. J. Mater. Chem. C, 2016, Supporting Information. '
    'https://www.rsc.org/suppdata/c6/tc/c6tc02660e/c6tc02660e1.pdf',
]
e = P[176]._p
for r in REFS:
    e = add_after(e, r, P[176])

# 참고문헌 번호를 본문 첫 인용 순서로 다시 매긴다 (원래 [4]가 5.4절에서 처음 나오므로)
MAP = {1: 1, 2: 2, 3: 3, 4: 9, 5: 4, 6: 5, 7: 6, 8: 7, 9: 8}
cite = re.compile(r'\[(\d+(?:,\s*\d+)*)\]')
body_end = P[172]._p   # '7. 참고문헌' 제목
for p in d.paragraphs:
    if p._p is body_end:
        break
    t = p.text
    if cite.search(t) and p.runs and not p._p.xpath('.//w:drawing'):
        set_text(p, cite.sub(lambda m: '[' + ', '.join(str(MAP[int(x)]) for x in re.split(r',\s*', m.group(1))) + ']', t))
ref_ps = []
started = False
for p in d.paragraphs:
    if p._p is body_end:
        started = True; continue
    if started and re.match(r'\[\d+\]', p.text):
        ref_ps.append(p)
texts = {MAP[int(re.match(r'\[(\d+)\]', p.text).group(1))]: re.sub(r'^\[\d+\]', '', p.text) for p in ref_ps}
assert sorted(texts) == list(range(1, 10)), sorted(texts)
for k, p in enumerate(ref_ps, 1):
    set_text(p, f'[{k}]' + texts[k])

d.save(dst)
print('saved', dst)
