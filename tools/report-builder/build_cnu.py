"""전남대 레포트 양식(.docx)에 content.js 본문을 넣어 제출 보고서를 만든다.

python3 build_cnu.py <content.js> <양식.docx> <out.docx> [--pages pages.json]

- 양식의 표지(로고·제목 상자·정보 표), 머리글 로고, 쪽 번호, 목차 필드는 그대로 쓴다.
- 양식의 안내 문구(파란 글씨, 확인사항)와 예시 본문은 지운다.
- 목차의 쪽 번호는 --pages(제목→쪽 번호)를 주면 채운다. 없으면 빈칸. (render 단계에서 PDF로 구해 다시 빌드)
- 결과 옆에 out.labels.json(그림·표 번호), out.headings.json(목차 항목)을 쓴다.
"""
import copy, json, os, re, subprocess, sys
from docx import Document
from docx.shared import Cm, Emu
from docx.oxml.ns import qn
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NSD = {'w': W}
TEXT_W = 9144  # 양식 본문 폭 (11904 - 1380*2 twips)


def E(tag, attrs=None, *children):
    el = etree.SubElement(etree.Element('x'), qn(tag)) if False else etree.Element(qn(tag))
    for k, v in (attrs or {}).items():
        el.set(qn(k), str(v))
    for c in children:
        if c is not None:
            el.append(c)
    return el


# ---------------------------------------------------------------- 내용 읽기
def load_content(path):
    js = f"const c=require({json.dumps(os.path.abspath(path))});process.stdout.write(JSON.stringify(c));"
    return json.loads(subprocess.run(['node', '-e', js], capture_output=True, text=True, check=True).stdout)


def number(C):
    """제목 번호, 그림·표 번호, 참고문헌 번호(처음 인용 순서)를 매긴다."""
    h = [0, 0, 0]; fig = tab = 0; labels = {}
    for b in C['body']:
        if b['t'] == 'h1':
            h = [h[0] + 1, 0, 0]; fig = tab = 0; b['n'] = f'{h[0]}.'
        elif b['t'] == 'h2':
            h[1] += 1; h[2] = 0; b['n'] = f'{h[0]}.{h[1]}'
        elif b['t'] == 'h3':
            h[2] += 1; b['n'] = f'{h[0]}.{h[1]}.{h[2]}'
        elif b['t'] in ('fig', 'photos'):
            fig += 1; b['label'] = f'그림 {h[0]}-{fig}'; labels['f:' + b['key']] = b['label']
        elif b['t'] == 'table':
            tab += 1; b['label'] = f'표 {h[0]}-{tab}'; labels['t:' + b['key']] = b['label']
    refs_h1 = h[0] + 1
    # 참고문헌: 본문 등장 순서
    order = []
    def scan(s):
        for k in re.findall(r'\[([A-Z]\d+)\]', s or ''):
            if k not in order:
                order.append(k)
    for b in C['body']:
        for s in texts_of(b):
            scan(s)
    missing = [k for k in order if k not in C['refs']]
    unused = [k for k in C['refs'] if k not in order]
    if missing or unused:
        raise SystemExit(f'참고문헌 키 오류 — 정의 안 됨: {missing}, 인용 안 됨: {unused}')
    refno = {k: i + 1 for i, k in enumerate(order)}

    def sub(s):
        s = re.sub(r'\{([ft]):(\w+)\}', lambda m: labels[m.group(1) + ':' + m.group(2)], s)
        return re.sub(r'\[([A-Z]\d+)\]', lambda m: f'[{refno[m.group(1)]}]', s)
    for b in C['body']:
        for k in ('x', 'cap'):
            if k in b:
                b[k] = sub(b[k])
        if 'items' in b and b['t'] == 'ol':
            b['items'] = [sub(x) for x in b['items']]
        if b['t'] == 'table':
            b['rows'] = [[sub(c) for c in r] for r in b['rows']]
    C['refs_list'] = [C['refs'][k] for k in order]
    C['refs_n'] = refs_h1
    return labels


def texts_of(b):
    if b['t'] in ('p', 'h1', 'h2', 'h3'):
        return [b['x']]
    if b['t'] == 'ol':
        return b['items']
    if b['t'] in ('fig', 'photos'):
        return [b.get('cap', '')]
    if b['t'] == 'table':
        return [b['cap']] + [c for r in b['rows'] for c in r]
    return []


# ---------------------------------------------------------------- 문단·글자
def rpr(sz=None, bold=False, color=None):
    r = E('w:rPr')
    if bold:
        r.append(E('w:b')); r.append(E('w:bCs'))
    if color:
        r.append(E('w:color', {'w:val': color}))
    if sz:
        r.append(E('w:sz', {'w:val': sz})); r.append(E('w:szCs', {'w:val': sz}))
    return r


def run(text, sz=None, bold=False, color=None):
    r = E('w:r', None, rpr(sz, bold, color))
    parts = text.split('\t')
    for i, part in enumerate(parts):
        if i:
            r.append(E('w:tab'))
        t = E('w:t', {'xml:space': 'preserve'}) if False else E('w:t')
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); t.text = part
        r.append(t)
    return r


def ppr(style=None, jc=None, before=None, after=None, line=None, ind=None, keepNext=False, keepLines=False,
        pageBreakBefore=False, tabs=None):
    p = E('w:pPr')
    if style:
        p.append(E('w:pStyle', {'w:val': style}))
    if keepNext:
        p.append(E('w:keepNext'))
    if keepLines:
        p.append(E('w:keepLines'))
    if pageBreakBefore:
        p.append(E('w:pageBreakBefore'))
    if tabs:
        t = E('w:tabs')
        for kind, pos, leader in tabs:
            a = {'w:val': kind, 'w:pos': pos}
            if leader:
                a['w:leader'] = leader
            t.append(E('w:tab', a))
        p.append(t)
    sp = {}
    if before is not None: sp['w:before'] = before
    if after is not None: sp['w:after'] = after
    if line is not None: sp['w:line'] = line; sp['w:lineRule'] = 'auto'
    if sp:
        p.append(E('w:spacing', sp))
    if ind:
        p.append(E('w:ind', ind))
    if jc:
        p.append(E('w:jc', {'w:val': jc}))
    return p


def para(runs, **kw):
    p = E('w:p', None, ppr(**kw))
    for r in runs:
        p.append(r)
    return p


BODY = dict(jc='both', after=100, line=300, ind={'w:firstLine': 200})
CAP_SZ = 19


# ---------------------------------------------------------------- 그림
class Pics:
    def __init__(self, doc):
        self.doc = doc

    def inline(self, path, width_emu, height_emu):
        part = self.doc.part
        rid, img = part.get_or_add_image(path)
        shape = part.new_pic_inline(path, Emu(width_emu), Emu(height_emu)) if False else None
        from docx.oxml.shape import CT_Inline
        cx, cy = int(width_emu), int(height_emu)
        inl = CT_Inline.new_pic_inline(next_id(), rid, os.path.basename(path), cx, cy)
        r = E('w:r')
        d = E('w:drawing'); d.append(inl); r.append(d)
        return r


_id = [100]
def next_id():
    _id[0] += 1
    return _id[0]


def img_size(path):
    from PIL import Image
    with Image.open(path) as im:
        return im.size


EMU_CM = 360000


# ---------------------------------------------------------------- 블록 → XML
def build_blocks(C, doc, base, pics):
    out = []
    body = C['body']
    first_h1 = True
    for i, b in enumerate(body):
        nxt = body[i + 1] if i + 1 < len(body) else None
        t = b['t']
        if t == 'h1':
            out.append(heading(b, 1, pageBreakBefore=first_h1)); first_h1 = False
        elif t == 'h2':
            out.append(heading(b, 2))
        elif t == 'h3':
            out.append(heading(b, 3))
        elif t == 'p':
            out.append(para([run(b['x'])], keepNext=bool(nxt and nxt['t'] == 'table'), **BODY))
        elif t == 'ol':
            for k, it in enumerate(b['items']):
                last = k == len(b['items']) - 1
                out.append(para([run(f'{k + 1}.\t{it}')], jc='both', after=120 if last else 30, line=290,
                                ind={'w:left': 600, 'w:hanging': 340}, tabs=[('left', 600, None)]))
        elif t == 'fig':
            path = os.path.join(base, b['img']); w, h = img_size(path)
            cx = min(b.get('w', 600), 610) / 96 * 2.54 * EMU_CM
            cy = cx * h / w
            out.append(para([pics.inline(path, cx, cy)], jc='center', before=120, after=40, keepNext=True))
            out.append(para([run(f"{b['label']}. {b['cap']}", sz=CAP_SZ)], jc='center', after=200, keepLines=True))
        elif t == 'photos':
            out.extend(photos(b, base, pics))
        elif t == 'table':
            out.append(para([run(f"{b['label']}. {b['cap']}", sz=CAP_SZ)], jc='center', before=160, after=80, keepNext=True))
            out.append(table(b))
            out.append(para([], after=60, line=240))
    return out


HEAD_SZ = {1: 28, 2: 24, 3: 22}
_bm = [500]


def heading(b, lv, pageBreakBefore=False):
    style = {1: '1', 2: 'NBSLH2', 3: 'NBSLH3'}[lv]
    before = {1: 360, 2: 240, 3: 160}[lv]
    p = para([], style=style, before=before, after=100 if lv < 3 else 60, keepNext=True, pageBreakBefore=pageBreakBefore,
             ind={'w:left': 0, 'w:firstLine': 0})
    if lv <= 2:
        _bm[0] += 1; b['bm'] = f'_Toc9{_bm[0]:05d}'
        p.append(E('w:bookmarkStart', {'w:id': _bm[0], 'w:name': b['bm']}))
    p.append(run(f"{b['n']} {b['x']}", sz=HEAD_SZ[lv], bold=True))
    if lv <= 2:
        p.append(E('w:bookmarkEnd', {'w:id': _bm[0]}))
    return p


def no_borders():
    b = E('w:tblBorders')
    for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b.append(E('w:' + s, {'w:val': 'nil'}))
    return b


def photos(b, base, pics):
    items = b['items']; H = b.get('h', 6.0) * EMU_CM
    sizes = []
    for it in items:
        w, h = img_size(os.path.join(base, it['img'])); sizes.append(H * w / h)
    gap = 0.25 * EMU_CM
    maxw = 15.0 * EMU_CM
    scale = min(1.0, (maxw - gap * (len(items) - 1)) / sum(sizes))
    out = []
    if len(items) == 1 and not items[0].get('sub'):
        it = items[0]
        out.append(para([pics.inline(os.path.join(base, it['img']), sizes[0] * scale, H * scale)], jc='center',
                        before=120, after=40, keepNext=True))
    else:
        # 사진 줄(안쪽 표) + 캡션을 칸 하나짜리 바깥 표에 넣고 행 나눔을 막아, 쪽이 바뀌어도 사진과 캡션이 같이 움직이게 한다
        tw = [int(w * scale / EMU_CM * 567) + 120 for w in sizes]
        inner = E('w:tbl')
        inner.append(E('w:tblPr', None, E('w:tblW', {'w:w': sum(tw), 'w:type': 'dxa'}), E('w:jc', {'w:val': 'center'}),
                       no_borders(), E('w:tblLayout', {'w:type': 'fixed'}),
                       E('w:tblCellMar', None, E('w:left', {'w:w': 60, 'w:type': 'dxa'}), E('w:right', {'w:w': 60, 'w:type': 'dxa'}))))
        g = E('w:tblGrid')
        for x in tw:
            g.append(E('w:gridCol', {'w:w': x}))
        inner.append(g)
        tr = E('w:tr', None, E('w:trPr', None, E('w:cantSplit')))
        for it, w, x in zip(items, sizes, tw):
            tc = E('w:tc', None, E('w:tcPr', None, E('w:tcW', {'w:w': x, 'w:type': 'dxa'}), E('w:vAlign', {'w:val': 'bottom'})))
            tc.append(para([pics.inline(os.path.join(base, it['img']), w * scale, H * scale)], jc='center', after=20, line=240))
            tc.append(para([run(it.get('sub', ''), sz=17)], jc='center', after=0, line=240))
            tr.append(tc)
        inner.append(tr)
        out.append(wrap_keep([inner, para([run(f"{b['label']}. {b['cap']}", sz=CAP_SZ)], jc='center', before=60, after=0)]))
        out.append(para([], after=120, line=240))
        return out
    out.append(para([run(f"{b['label']}. {b['cap']}", sz=CAP_SZ)], jc='center', before=60, after=200, keepLines=True))
    return out


def wrap_keep(children):
    """칸 하나·행 나눔 금지 표로 감싼다 (테두리 없음)."""
    t = E('w:tbl')
    t.append(E('w:tblPr', None, E('w:tblW', {'w:w': TEXT_W, 'w:type': 'dxa'}), E('w:jc', {'w:val': 'center'}), no_borders(),
               E('w:tblLayout', {'w:type': 'fixed'}),
               E('w:tblCellMar', None, E('w:left', {'w:w': 0, 'w:type': 'dxa'}), E('w:right', {'w:w': 0, 'w:type': 'dxa'}))))
    t.append(E('w:tblGrid', None, E('w:gridCol', {'w:w': TEXT_W})))
    tc = E('w:tc', None, E('w:tcPr', None, E('w:tcW', {'w:w': TEXT_W, 'w:type': 'dxa'})))
    for c in children:
        tc.append(c)
    t.append(E('w:tr', None, E('w:trPr', None, E('w:cantSplit')), tc))
    return t


def table(b):
    cols = b['cols']; assert sum(cols) <= TEXT_W, (b['key'], sum(cols))
    tbl = E('w:tbl')
    tbl.append(E('w:tblPr', None, E('w:tblStyle', {'w:val': 'a3'}), E('w:tblW', {'w:w': sum(cols), 'w:type': 'dxa'}),
                 E('w:jc', {'w:val': 'center'}), E('w:tblLayout', {'w:type': 'fixed'}),
                 E('w:tblCellMar', None, E('w:top', {'w:w': 40, 'w:type': 'dxa'}), E('w:left', {'w:w': 90, 'w:type': 'dxa'}),
                   E('w:bottom', {'w:w': 40, 'w:type': 'dxa'}), E('w:right', {'w:w': 90, 'w:type': 'dxa'}))))
    g = E('w:tblGrid')
    for c in cols:
        g.append(E('w:gridCol', {'w:w': c}))
    tbl.append(g)
    rows = [b['head']] + b['rows']
    for ri, r in enumerate(rows):
        head = ri == 0
        trp = E('w:trPr', None, E('w:cantSplit'))
        if head:
            trp.append(E('w:tblHeader'))
        tr = E('w:tr', None, trp)
        for ci, (txt, wdt) in enumerate(zip(r, cols)):
            tcp = E('w:tcPr', None, E('w:tcW', {'w:w': wdt, 'w:type': 'dxa'}), E('w:vAlign', {'w:val': 'center'}))
            if head:
                tcp.insert(1, E('w:shd', {'w:val': 'clear', 'w:color': 'auto', 'w:fill': 'E7EEF6'}))
            tc = E('w:tc', None, tcp)
            tc.append(para([run(txt, sz=CAP_SZ, bold=head or ci == 0 and len(cols) > 2 and False)],
                           jc='center' if head or ci == 0 else 'left', after=0, line=260,
                           keepNext=ri < len(rows) - 1))
            tr.append(tc)
        tbl.append(tr)
    return tbl


# ---------------------------------------------------------------- 양식 손보기
def add_heading_styles(doc):
    st = doc.styles.element
    for sid, name, lvl, sz in (('NBSLH2', 'heading 2', 1, 24), ('NBSLH3', 'heading 3', 2, 22)):
        if st.find(f'w:style[@w:styleId="{sid}"]', NSD) is not None:
            continue
        s = E('w:style', {'w:type': 'paragraph', 'w:styleId': sid}, E('w:name', {'w:val': name}),
              E('w:basedOn', {'w:val': 'a'}), E('w:next', {'w:val': 'a'}), E('w:uiPriority', {'w:val': 9}),
              E('w:unhideWhenUsed'), E('w:qFormat'),
              E('w:pPr', None, E('w:keepNext'), E('w:outlineLvl', {'w:val': lvl})),
              E('w:rPr', None, E('w:rFonts', {'w:asciiTheme': 'majorHAnsi', 'w:eastAsiaTheme': 'majorEastAsia', 'w:hAnsiTheme': 'majorHAnsi'}),
                E('w:b'), E('w:sz', {'w:val': sz}), E('w:szCs', {'w:val': sz})))
        st.append(s)


def set_cover(doc, C):
    body = doc.element.body
    title = C['cover']['title']
    done = 0
    for tx in body.iter(qn('w:txbxContent')):
        txt = ''.join(tx.itertext())
        if '제목' not in txt:
            continue
        p = tx.find('w:p', NSD)
        runs = p.findall('w:r', NSD)
        keep = copy.deepcopy(runs[1].find('w:rPr', NSD))  # 굵은 밑줄 제목 글자 서식
        for r in p.findall('w:r', NSD):
            p.remove(r)
        for x in p.findall('w:bookmarkStart', NSD) + p.findall('w:bookmarkEnd', NSD) + p.findall('w:proofErr', NSD):
            p.remove(x)
        sz = keep.find('w:sz', NSD); szc = keep.find('w:szCs', NSD)
        fs = C['cover'].get('title_sz', 52)
        sz.set(qn('w:val'), str(fs)); szc.set(qn('w:val'), str(fs))
        u = keep.find('w:u', NSD)
        if u is not None:
            keep.remove(u)
        r = E('w:r'); r.append(keep); t = E('w:t'); t.text = title; r.append(t); p.append(r)
        for extra in tx.findall('w:p', NSD)[1:]:
            tx.remove(extra)
        done += 1
    assert done >= 1, '표지 제목 상자를 찾지 못함'
    # 정보 표
    tbl = next(t for t in body.iter(qn('w:tbl')) if '담당교수' in ''.join(t.itertext()))
    for tr in tbl.findall('w:tr', NSD):
        tcs = tr.findall('w:tc', NSD)
        key = re.sub(r'[\s•:]', '', ''.join(x.text or '' for x in tcs[0].iter(qn('w:t'))))
        val = C['cover']['info'].get(key, '')
        if val:
            p = tcs[1].find('w:p', NSD)
            jc = p.find('w:pPr/w:jc', NSD)
            if jc is not None:
                jc.set(qn('w:val'), 'left')
            r = E('w:r', None, E('w:rPr', None, E('w:rFonts', {'w:ascii': '맑은 고딕', 'w:eastAsia': '맑은 고딕', 'w:hAnsi': '맑은 고딕'}),
                                   E('w:sz', {'w:val': 26}), E('w:szCs', {'w:val': 26})))
            t = E('w:t'); t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); t.text = ' ' + val; r.append(t); p.append(r)


def strip_template_body(doc):
    body = doc.element.body
    sdt = body.find('w:sdt', NSD)
    sect = body.find('w:sectPr', NSD)
    el = sdt.getnext()
    while el is not None and el is not sect:
        nx = el.getnext(); body.remove(el); el = nx
    return sdt, sect


def build_toc(sdt, entries, pages):
    content = sdt.find('w:sdtContent', NSD)
    ps = content.findall('w:p', NSD)
    for p in ps[1:]:
        content.remove(p)
    def fld(kind):
        return E('w:r', None, rpr(), E('w:fldChar', {'w:fldCharType': kind}))
    def instr(s):
        r = E('w:r', None, rpr()); it = E('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); it.text = s; r.append(it); return r
    for i, (lv, num, title, bm) in enumerate(entries):
        sz = 24 if lv == 1 else 22
        if lv == 1:
            p = para([], style='10', before=160 if i else 240, after=60, tabs=[('left', 520, None), ('right', 9100, 'dot')])
        else:
            p = para([], style='2', before=0, after=40, ind={'w:left': 300}, tabs=[('left', 1000, None), ('right', 9100, 'dot')])
        if i == 0:
            p.append(fld('begin')); p.append(instr(' TOC \\o "1-2" \\h \\z \\u ')); p.append(fld('separate'))
        h = E('w:hyperlink', {'w:anchor': bm, 'w:history': 1})
        h.append(run(f'{num}\t{title}\t', sz=sz, bold=lv == 1))
        h.append(fld('begin')); h.append(instr(f' PAGEREF {bm} \\h ')); h.append(fld('separate'))
        h.append(run(str(pages.get(bm, '')), sz=sz, bold=lv == 1)); h.append(fld('end'))
        p.append(h)
        content.append(p)
    content.append(para([fld('end')], after=0))


def main():
    a = sys.argv[1:]
    content_path, tpl, out = a[0], a[1], a[2]
    pages = json.load(open(a[a.index('--pages') + 1])) if '--pages' in a else {}
    C = load_content(content_path)
    base = os.path.dirname(os.path.abspath(content_path))
    labels = number(C)
    doc = Document(tpl)
    add_heading_styles(doc)
    set_cover(doc, C)
    sdt, sect = strip_template_body(doc)
    pics = Pics(doc)
    blocks = build_blocks(C, doc, base, pics)
    # 참고문헌
    rb = {'t': 'h1', 'n': f"{C['refs_n']}.", 'x': '참고문헌'}
    blocks.append(heading(rb, 1))
    for i, r in enumerate(C['refs_list']):
        blocks.append(para([run(f'[{i + 1}]\t{r}', sz=CAP_SZ)], jc='left', after=80, line=260,
                           ind={'w:left': 520, 'w:hanging': 520}, tabs=[('left', 520, None)]))
    for el in blocks:
        sect.addprevious(el)
    entries = [(1 if b['t'] == 'h1' else 2, b['n'], b['x'], b['bm']) for b in C['body'] if b['t'] in ('h1', 'h2')]
    entries.append((1, rb['n'], rb['x'], rb['bm']))
    build_toc(sdt, entries, pages)
    # 문서 속성
    cp = doc.core_properties
    cp.title = C['cover']['title']; cp.author = ''; cp.last_modified_by = ''; cp.comments = ''
    doc.save(out)
    stem = re.sub(r'\.docx$', '', out)
    json.dump({k.split(':', 1)[1]: v for k, v in labels.items()}, open(stem + '.labels.json', 'w'), ensure_ascii=False, indent=1)
    json.dump([{'lv': e[0], 'text': f'{e[1]} {e[2]}', 'bm': e[3]} for e in entries], open(stem + '.headings.json', 'w'),
              ensure_ascii=False, indent=1)
    print('ok', out, '참고문헌', len(C['refs_list']))


if __name__ == '__main__':
    main()
